"""Build the distributed Office templates (development only: python-docx, openpyxl)."""
import csv
import json
from pathlib import Path
import zipfile
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.page import PageMargins

ROOT = Path(__file__).resolve().parents[1]
FOLDER = ROOT / 'templates'
TEAL = '166A58'


def word(source, target, label):
    doc = Document()
    section = doc.sections[0]
    section.page_height, section.page_width = Inches(11.69), Inches(8.27)
    section.top_margin = section.bottom_margin = Inches(.65)
    normal = doc.styles['Normal']
    normal.font.name, normal.font.size = 'Calibri', Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    for name in ('Title', 'Heading 1', 'Heading 2'):
        doc.styles[name].font.color.rgb = RGBColor.from_string(TEAL)
    section.header.paragraphs[0].text = 'SHEguru  |  Safety, made manageable.'
    doc.add_heading(label, 0)
    doc.add_paragraph('WORKING TEMPLATE — complete, adapt and approve for your workplace.', 'Subtitle')
    table = None
    for line in source.read_text().splitlines()[1:]:
        if not line.strip():
            table = None
            continue
        if line.isupper() and len(line) < 115:
            doc.add_heading(line, 1)
            table = None
        elif line.rstrip().endswith(':') and len(line) < 100:
            if table is None:
                table = doc.add_table(rows=0, cols=2)
                table.style = 'Light Shading Accent 1'
            cells = table.add_row().cells
            cells[0].text = line.rstrip(':')
            cells[1].text = '\n'
            cells[0].width, cells[1].width = Inches(2.5), Inches(4)
        else:
            table = None
            doc.add_paragraph(line)
    footer = section.footer.paragraphs[0]
    footer.text = 'SHEguru • Draft for workplace review  |  Page '
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    doc.save(target)


def workbook(source, target, label):
    with source.open(newline='') as f:
        rows = list(csv.DictReader(f))
    wb = Workbook()
    sheet = wb.active
    sheet.title = 'Risk assessment' if source.stem == 'risk-assessment' else 'Inspection'
    is_risk = source.stem == 'risk-assessment'
    headings = (['Task / activity', 'Hazard / people exposed', 'Existing controls', 'Initial risk / method', 'Further controls', 'Action owner / due date', 'Residual risk', 'Approval / review'] if is_risk else
                ['Check item', 'What to check', 'Finding', 'Defect / corrective action', 'Action owner', 'Due date', 'Completed / verified', 'Release / next check'])
    sheet.merge_cells('A1:H1')
    sheet['A1'] = 'SHEguru | ' + label
    sheet['A1'].font = Font(name='Calibri', size=20, bold=True, color='FFFFFF')
    sheet['A1'].fill = PatternFill('solid', fgColor=TEAL)
    sheet.row_dimensions[1].height = 38
    fields = [('Company:', 'Workplace / area:'), ('Assessor / inspector:', 'Date / time:'), ('Activity / asset ID:', 'Review / next inspection:')]
    for r, (left, right) in enumerate(fields, 3):
        sheet.cell(r, 1, left).font = Font(bold=True)
        sheet.cell(r, 1).alignment = Alignment(wrap_text=True, vertical='center')
        sheet.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        sheet.cell(r, 5, right).font = Font(bold=True)
        sheet.cell(r, 5).alignment = Alignment(wrap_text=True, vertical='center')
        sheet.merge_cells(start_row=r, start_column=6, end_row=r, end_column=8)
        sheet.row_dimensions[r].height = 26
    sheet.merge_cells('A7:H7')
    sheet['A7'] = ('Assess the actual task and record your risk method. Complete controls, owners and review before approval.' if is_risk else
                   'Adapt to the equipment and manufacturer instructions. Record findings; isolate defects and authorise return to use only when safe.')
    sheet['A7'].alignment = Alignment(wrap_text=True, vertical='center')
    sheet.row_dimensions[7].height = 36
    for c, title in enumerate(headings, 1):
        cell = sheet.cell(9, c, title)
        cell.fill = PatternFill('solid', fgColor=TEAL)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    sheet.row_dimensions[9].height = 32
    if is_risk:
        entries = [[''] * 8 for _ in range(12)]
    elif source.stem == 'inspection-sheet':
        entries = [['', 'Enter a workplace-specific inspection item and applicable check.', '', '', '', '', '', ''] for _ in range(12)]
    else:
        entries = [[r['Check item'], r['Inspection criterion'], '', '', '', '', '', ''] for r in rows]
    for n, data in enumerate(entries, 10):
        sheet.append(data)
        # Explicit row numbers keep the form stable even with merged header ranges.
        for c, value in enumerate(data, 1):
            cell = sheet.cell(n, c, value)
            cell.alignment = Alignment(wrap_text=True, vertical='top')
            cell.border = Border(bottom=Side(style='thin', color='DCE5E5'))
            if n % 2 == 0:
                cell.fill = PatternFill('solid', fgColor='F0F6F4')
        sheet.row_dimensions[n].height = 75
    end = 9 + len(entries)
    if not is_risk:
        validation = DataValidation(type='list', formula1='"OK,Defect,Not applicable"', allow_blank=True)
        validation.errorTitle, validation.error = 'Choose a finding', 'Choose OK, Defect or Not applicable.'
        validation.showErrorMessage = True
        sheet.add_data_validation(validation)
        validation.add(f'C10:C{end}')
    widths = [28, 58, 20, 42, 22, 18, 30, 30]
    if is_risk:
        widths = [30, 40, 36, 25, 40, 28, 25, 28]
    for letter, width in zip('ABCDEFGH', widths):
        sheet.column_dimensions[letter].width = width
    sheet.freeze_panes = 'C10'
    sheet.auto_filter.ref = f'A9:H{end}'
    sheet.sheet_view.zoomScale = 80
    sheet.print_options.horizontalCentered = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = '1:9'
    sheet.print_area = f'A1:H{end}'
    sheet.page_margins = PageMargins(left=.25, right=.25, top=.4, bottom=.4)
    sheet.oddFooter.center.text = 'SHEguru | Page &P of &N'
    guide = wb.create_sheet('Read me')
    notes = [label, 'Blank working template — adapt and approve for the workplace.',
             'Fill in company, workplace, date, responsible person and the asset / task details.',
             'Blue / green headers are labels. Enter your answers in the blank cells.',
             'Wrap Text is enabled. Increase row heights if you enter longer notes.',
             'Blank fields do not mean satisfactory. Record the actual finding and evidence.',
             'Save a completed copy with a new name, then upload it to SHEguru Legal documents.',
             'Downloading this form does not complete a requirement or approve equipment for use.',
             'Apply the manufacturer instructions and the inspection / competence requirements for the actual equipment.',
             'Inspection sheets are pre-use / workplace aids and do not replace specialist examinations or statutory tests.']
    if is_risk:
        notes += ['Define and record your likelihood and severity scales and risk-rating method.',
                  'Describe persons exposed, actual controls and additional controls using the hierarchy of controls.',
                  'Record residual risk after controls and obtain the appropriate review / approval.']
    for i, text in enumerate(notes, 1):
        guide.cell(i, 1, text).alignment = Alignment(wrap_text=True, vertical='top')
        guide.row_dimensions[i].height = 36
    guide.column_dimensions['A'].width = 110
    guide['A1'].font = Font(size=18, bold=True, color=TEAL)
    wb.save(target)


def main():
    catalogue = json.loads((FOLDER / 'catalogue.json').read_text())
    for row in catalogue:
        # Keep editable text / CSV sources for future regeneration; distribute Office files.
        source = FOLDER / (Path(row['file']).stem + ('.csv' if Path(row['file']).suffix in ('.csv', '.xlsx') else '.txt'))
        target = source.with_suffix('.xlsx' if source.suffix == '.csv' else '.docx')
        (workbook if source.suffix == '.csv' else word)(source, target, row['label'])
        row['file'] = target.name
    (FOLDER / 'catalogue.json').write_text(json.dumps(catalogue, indent=2) + '\n')
    with zipfile.ZipFile(ROOT / 'SHEguru-Word-Excel-templates.zip', 'w', zipfile.ZIP_DEFLATED) as pack:
        for row in catalogue:
            section = '1-Appointments' if row['section'] == 'appointment' else '2-Legal-documents'
            pack.write(FOLDER / row['file'], section + '/' + row['file'])
        pack.writestr('START-HERE.txt', 'SHEguru templates\n\nOpen appointment letters and reports (.docx) in Word.\nOpen risk assessments and inspection sheets (.xlsx) in Excel.\nComplete and review for your workplace, then save a separate copy.\nThe section 37(2) contractor agreement is a draft for appropriate review.\nThese files do not update the SHEguru app or change your records.\n')


if __name__ == '__main__':
    main()
